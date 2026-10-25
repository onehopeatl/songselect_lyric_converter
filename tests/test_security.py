import io
import os
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import sys
import subprocess
import tempfile
import unittest
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'app'))
from pptx import Presentation
import webapp
from songselect_to_pptx import lyric_converter


def lyrics(title='Example Song', line='An original test lyric'):
    return f'{title}\n\nVerse 1\n{line}\n\nTest Author\n© 2026 Test\nCCLI Song # 123\n'.encode()


class SecurityTests(unittest.TestCase):
    def upload(self, content=None, filename='song.txt'):
        with webapp.app.test_client() as client:
            return client.post('/', data={'files[]': (io.BytesIO(content if content is not None else lyrics()), filename)})

    def test_conversion_preserves_lyrics_and_license(self):
        with patch.dict(os.environ, {'CCLI_LIC_NUM': 'TEST-123'}):
            response = self.upload()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, 'application/zip')
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            self.assertEqual(archive.namelist(), ['Example_Song.pptx'])
            prs = Presentation(io.BytesIO(archive.read(archive.namelist()[0])))
        self.assertEqual(len(prs.slides), 2)
        self.assertIn('TEST-123', prs.slides[0].shapes[1].text)
        self.assertEqual(prs.slides[1].shapes[0].text, 'An original test lyric')
        response.close()

    def test_titles_cannot_escape_output_directory(self):
        with tempfile.TemporaryDirectory() as work:
            root = Path(work)
            input_path = root / 'input.txt'
            output_dir = root / 'output'
            output_dir.mkdir()
            for title in ['../../escaped', '/tmp/escaped', r'..\..\escaped', '🎵', 'a' * 1000]:
                input_path.write_bytes(lyrics(title))
                result = Path(lyric_converter(input_path, output_dir))
                self.assertEqual(result.parent.resolve(), output_dir.resolve())
                self.assertLessEqual(len(result.name), 155)
            self.assertFalse((root / 'escaped.pptx').exists())

    def test_duplicate_titles_are_not_overwritten(self):
        with webapp.app.test_client() as client:
            response = client.post('/', data={'files[]': [
                (io.BytesIO(lyrics(line='First lyric')), 'same.txt'),
                (io.BytesIO(lyrics(line='Second lyric')), 'same.txt'),
            ]})
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            self.assertEqual(archive.namelist(), ['Example_Song.pptx', 'Example_Song-2.pptx'])
            texts = [Presentation(io.BytesIO(archive.read(name))).slides[1].shapes[0].text for name in archive.namelist()]
        self.assertEqual(texts, ['First lyric', 'Second lyric'])
        response.close()

    def test_concurrent_requests_are_isolated_and_cleaned(self):
        with tempfile.TemporaryDirectory() as work, patch.object(tempfile, 'tempdir', work):
            def convert(number):
                response = self.upload(lyrics(line=f'Unique lyric {number}'))
                self.assertEqual(response.status_code, 200)
                with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
                    text = Presentation(io.BytesIO(archive.read(archive.namelist()[0]))).slides[1].shapes[0].text
                response.close()
                return text
            with ThreadPoolExecutor(max_workers=4) as executor:
                self.assertEqual(list(executor.map(convert, range(8))), [f'Unique lyric {n}' for n in range(8)])
            self.assertEqual(list(Path(work).iterdir()), [])

    def test_invalid_uploads_have_safe_visible_errors(self):
        for content, filename in [(b'', 'empty.txt'), (b'\xff', 'bad.txt'), (b'bad', 'bad.html'), (lyrics(), '')]:
            response = self.upload(content, filename)
            self.assertEqual(response.status_code, 400)
            self.assertIn(b'role="alert"', response.data)
            response.close()
        with self.assertLogs(webapp.app.logger, level='ERROR'), patch.object(webapp, 'lyric_converter', side_effect=RuntimeError('/private/secret')):
            response = self.upload()
        self.assertEqual(response.status_code, 500)
        self.assertNotIn(b'/private/secret', response.data)
        response.close()

    def test_upload_resource_limits(self):
        response = self.upload(b'x' * (webapp.MAX_LYRIC_FILE_SIZE + 1))
        self.assertEqual(response.status_code, 413)
        response.close()
        response = self.upload(b'x' * (10 * 1024 * 1024 + 1))
        self.assertEqual(response.status_code, 413)
        response.close()
        with webapp.app.test_client() as client:
            response = client.post('/', data={'files[]': [(io.BytesIO(lyrics()), f'{n}.txt') for n in range(11)]})
        self.assertEqual(response.status_code, 400)
        response.close()
        response = self.upload(b'Song\n' + b'Verse 1\nTest\n' * 201)
        self.assertEqual(response.status_code, 400)
        response.close()

    def test_valid_large_lyric_file_is_accepted(self):
        # Multipart framing must not accidentally lower the documented file limit.
        payload = lyrics() + b'metadata\n' * 16000
        response = self.upload(payload)
        self.assertEqual(response.status_code, 200)
        response.close()

    def test_forwarded_headers_are_not_trusted_by_default(self):
        with webapp.app.test_request_context('/', headers={
            'X-Forwarded-For': '1.2.3.4', 'X-Forwarded-Host': 'evil.example',
            'X-Forwarded-Proto': 'https', 'X-Forwarded-Prefix': '/evil',
        }):
            self.assertEqual(webapp.request.host, 'localhost')
            self.assertEqual(webapp.request.scheme, 'http')
            self.assertEqual(webapp.request.script_root, '')

    def test_explicit_proxy_trust_preserves_host_and_prefix(self):
        source = """
import webapp
from flask import request
webapp.app.add_url_rule('/probe', 'probe', lambda: {
    'address': request.remote_addr, 'scheme': request.scheme,
    'host': request.host, 'prefix': request.script_root,
})
with webapp.app.test_client() as client:
    result = client.get('/probe', headers={
        'X-Forwarded-For': '198.51.100.66, 1.2.3.4', 'X-Forwarded-Proto': 'https',
        'X-Forwarded-Host': 'evil.example', 'X-Forwarded-Prefix': '/evil',
    }).json
assert result == {'address': '1.2.3.4', 'scheme': 'https', 'host': 'localhost', 'prefix': ''}, result
"""
        env = dict(os.environ, TRUSTED_PROXY_COUNT='1', PYTHONPATH=str(Path(webapp.__file__).parent))
        subprocess.run([sys.executable, '-c', source], env=env, check=True)

    def test_archive_is_closed_after_download_and_failures(self):
        original_spool = tempfile.SpooledTemporaryFile
        spools = []

        def track_spool(*args, **kwargs):
            spool = original_spool(*args, **kwargs)
            spools.append(spool)
            return spool

        with patch.object(webapp.tempfile, 'SpooledTemporaryFile', side_effect=track_spool):
            response = self.upload()
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.data)
            response.close()
            self.assertTrue(all(spool.closed for spool in spools))
            response = self.upload(b'')
            self.assertEqual(response.status_code, 400)
            response.close()
            self.assertTrue(all(spool.closed for spool in spools))

    def test_cross_site_uploads_rejected_and_headers_set(self):
        with webapp.app.test_client() as client:
            response = client.post('/', headers={'Sec-Fetch-Site': 'cross-site'})
            self.assertEqual(response.status_code, 403)
            response = client.get('/')
        self.assertEqual(response.headers['X-Content-Type-Options'], 'nosniff')
        self.assertEqual(response.headers['X-Frame-Options'], 'DENY')
        self.assertIn("script-src 'self'", response.headers['Content-Security-Policy'])
        self.assertIn(b'/static/app.js', response.data)

    def test_bom_and_whitespace_title(self):
        response = self.upload(b'\xef\xbb\xbf  ' + lyrics())
        self.assertEqual(response.status_code, 200)
        response.close()


if __name__ == '__main__':
    unittest.main()
